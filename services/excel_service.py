import io
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelService:
    @staticmethod
    def create_student_report(records: list, school_name: str, title_meta: dict):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Attendance"

        # Theme Colors
        teal_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        teal_sub_fill = PatternFill(start_color="115E59", end_color="115E59", fill_type="solid")
        header_fill = PatternFill(start_color="134E4A", end_color="134E4A", fill_type="solid")
        summary_hdr_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Calibri", size=10, bold=True, color="166534")
        
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        absent_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        white_title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        white_meta_font = Font(name="Calibri", size=10, color="E0F2FE")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title Block
        ws.merge_cells("A1:I1")
        ws["A1"] = school_name.upper()
        ws["A1"].font = white_title_font
        ws["A1"].fill = teal_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:I2")
        meta_str = f"STUDENT ATTENDANCE REPORT  |  {title_meta.get('filter_desc', '')}  |  Generated on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A2"] = meta_str
        ws["A2"].font = white_meta_font
        ws["A2"].fill = teal_sub_fill
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        if not records:
            ws.merge_cells("A4:I5")
            ws["A4"] = "No attendance records found."
            ws["A4"].font = Font(name="Calibri", size=14, bold=True, color="64748B")
            ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer

        # Detailed Table Header
        headers = [
            "Roll Number", "Student Name", "Class", "Subject", 
            "Date", "Day", "Period", "Status", "Teacher"
        ]
        start_row = 4
        for col_num, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num)
            cell.value = h
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[start_row].height = 24

        curr_row = start_row + 1
        summary_data = {} # studentName -> {roll, total, present, absent}

        for idx, rec in enumerate(records):
            student_name = rec.get("studentName", "Unknown")
            roll_no = rec.get("rollNumber", "-")
            status = rec.get("status", "Absent")

            # Day of week from date
            date_val = rec.get("date", "")
            day_name = "-"
            try:
                dt = datetime.datetime.strptime(date_val, "%Y-%m-%d")
                day_name = dt.strftime("%A")
            except Exception:
                pass

            row_data = [
                roll_no,
                student_name,
                rec.get("className", rec.get("classId", "-")),
                rec.get("subjectName", rec.get("subjectId", "-")),
                date_val,
                day_name,
                rec.get("period", "-"),
                status,
                rec.get("teacherName", "-")
            ]

            fill_to_use = zebra_fill if idx % 2 == 1 else white_fill

            for col_num, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col_num)
                cell.value = val
                cell.font = regular_font
                cell.fill = fill_to_use
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col_num in (1, 3, 5, 6, 7) else "left", vertical="center")

                # Custom highlight on status
                if col_num == 8:
                    if status.lower() == "present":
                        cell.fill = present_fill
                        cell.font = present_font
                    else:
                        cell.fill = absent_fill
                        cell.font = absent_font

            # Accumulate summary
            if student_name not in summary_data:
                summary_data[student_name] = {"roll": roll_no, "total": 0, "present": 0, "absent": 0}
            summary_data[student_name]["total"] += 1
            if status.lower() == "present":
                summary_data[student_name]["present"] += 1
            else:
                summary_data[student_name]["absent"] += 1

            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        # Summary Section
        curr_row += 2
        ws.merge_cells(f"A{curr_row}:E{curr_row}")
        ws[f"A{curr_row}"] = "ATTENDANCE SUMMARY BY STUDENT"
        ws[f"A{curr_row}"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        ws[f"A{curr_row}"].fill = summary_hdr_fill
        ws[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        sum_headers = ["Roll No", "Student Name", "Total Classes", "Present", "Absent", "Attendance %"]
        for col_num, sh in enumerate(sum_headers, 1):
            cell = ws.cell(row=curr_row, column=col_num)
            cell.value = sh
            cell.font = header_font
            cell.fill = teal_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

        for idx, (s_name, s_info) in enumerate(sorted(summary_data.items(), key=lambda x: str(x[1]["roll"]))):
            tot = s_info["total"]
            prs = s_info["present"]
            abs_cnt = s_info["absent"]
            pct = (prs / tot * 100) if tot > 0 else 0.0

            s_row = [
                s_info["roll"],
                s_name,
                tot,
                prs,
                abs_cnt,
                f"{pct:.1f}%"
            ]
            fill_to_use = zebra_fill if idx % 2 == 1 else white_fill
            for col_num, val in enumerate(s_row, 1):
                cell = ws.cell(row=curr_row, column=col_num)
                cell.value = val
                cell.font = bold_font if col_num in (2, 6) else regular_font
                cell.fill = fill_to_use
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col_num in (1, 3, 4, 5, 6) else "left", vertical="center")
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        # Adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and "\n" not in val_str and cell.row > 2:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def create_teacher_report(records: list, school_name: str, title_meta: dict):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Teacher Attendance"

        teal_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        teal_sub_fill = PatternFill(start_color="115E59", end_color="115E59", fill_type="solid")
        header_fill = PatternFill(start_color="134E4A", end_color="134E4A", fill_type="solid")
        summary_hdr_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Calibri", size=10, bold=True, color="166534")
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        absent_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        white_title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        white_meta_font = Font(name="Calibri", size=10, color="E0F2FE")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.merge_cells("A1:E1")
        ws["A1"] = school_name.upper()
        ws["A1"].font = white_title_font
        ws["A1"].fill = teal_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:E2")
        meta_str = f"TEACHER ATTENDANCE REPORT  |  {title_meta.get('filter_desc', '')}  |  Generated on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A2"] = meta_str
        ws["A2"].font = white_meta_font
        ws["A2"].fill = teal_sub_fill
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        if not records:
            ws.merge_cells("A4:E5")
            ws["A4"] = "No attendance records found."
            ws["A4"].font = Font(name="Calibri", size=14, bold=True, color="64748B")
            ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer

        headers = ["Teacher UID", "Teacher Name", "Date", "Day", "Status"]
        start_row = 4
        for col_num, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num)
            cell.value = h
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[start_row].height = 24

        curr_row = start_row + 1
        summary_data = {}

        for idx, rec in enumerate(records):
            t_uid = rec.get("teacherCode") or rec.get("teacherUid", "-")
            t_name = rec.get("teacherName", "Unknown")
            status = rec.get("status", "Absent")
            date_val = rec.get("date", "")
            day_name = "-"
            try:
                dt = datetime.datetime.strptime(date_val, "%Y-%m-%d")
                day_name = dt.strftime("%A")
            except Exception:
                pass

            row_data = [t_uid, t_name, date_val, day_name, status]
            fill_to_use = zebra_fill if idx % 2 == 1 else white_fill

            for col_num, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col_num)
                cell.value = val
                cell.font = regular_font
                cell.fill = fill_to_use
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col_num in (1, 3, 4) else "left", vertical="center")
                if col_num == 5:
                    if status.lower() == "present":
                        cell.fill = present_fill
                        cell.font = present_font
                    else:
                        cell.fill = absent_fill
                        cell.font = absent_font

            if t_name not in summary_data:
                summary_data[t_name] = {"uid": t_uid, "total": 0, "present": 0, "absent": 0}
            summary_data[t_name]["total"] += 1
            if status.lower() == "present":
                summary_data[t_name]["present"] += 1
            else:
                summary_data[t_name]["absent"] += 1

            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        # Summary Section
        curr_row += 2
        ws.merge_cells(f"A{curr_row}:E{curr_row}")
        ws[f"A{curr_row}"] = "TEACHER ATTENDANCE SUMMARY"
        ws[f"A{curr_row}"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        ws[f"A{curr_row}"].fill = summary_hdr_fill
        ws[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        sum_headers = ["Teacher UID", "Teacher Name", "Total Marked Days", "Present", "Absent", "Attendance %"]
        for col_num, sh in enumerate(sum_headers, 1):
            cell = ws.cell(row=curr_row, column=col_num)
            cell.value = sh
            cell.font = header_font
            cell.fill = teal_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

        for idx, (t_name, s_info) in enumerate(sorted(summary_data.items())):
            tot = s_info["total"]
            prs = s_info["present"]
            abs_cnt = s_info["absent"]
            pct = (prs / tot * 100) if tot > 0 else 0.0

            s_row = [
                s_info["uid"],
                t_name,
                tot,
                prs,
                abs_cnt,
                f"{pct:.1f}%"
            ]
            fill_to_use = zebra_fill if idx % 2 == 1 else white_fill
            for col_num, val in enumerate(s_row, 1):
                cell = ws.cell(row=curr_row, column=col_num)
                cell.value = val
                cell.font = bold_font if col_num in (2, 6) else regular_font
                cell.fill = fill_to_use
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col_num in (1, 3, 4, 5, 6) else "left", vertical="center")
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and "\n" not in val_str and cell.row > 2:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
